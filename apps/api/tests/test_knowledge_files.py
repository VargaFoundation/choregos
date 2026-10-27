"""Serving mode `files` (S33-02, ADR 0051): a run is served the documents pinned when it was prepared."""

from __future__ import annotations

from typing import Any

from choregos_contracts import knowledge_digest
from httpx import AsyncClient

from . import knowledge_fixtures as fx

BASE = "/api/v1/orgs/varga/knowledge"


async def _run(project: dict[str, Any], knowledge: list[dict[str, Any]]) -> str:
    from choregos_api.db.models import Run, WorkItem
    from choregos_api.db.session import session_scope
    from choregos_api.security import mint_run_token

    async with session_scope() as session:
        item = WorkItem(project_id=project["id"], tracker_key="BILLING-9", title="T", state="ready")
        session.add(item)
        await session.flush()
        session.add(
            Run(
                id="run-kb",
                work_item_id=item.id,
                project_id=project["id"],
                stage_role="implement",
                status="running",
                stage_input={"knowledge": knowledge},
            )
        )
    return mint_run_token("run-kb", project_slug="billing-api", work_item_key="BILLING-9", ttl_minutes=30)


async def _pinned(project: dict[str, Any]) -> list[dict[str, Any]]:
    from choregos_api.db.models import Organization
    from choregos_api.db.session import session_scope
    from choregos_api.knowledge.serving import pin
    from sqlalchemy import select

    async with session_scope() as session:
        org = (await session.execute(select(Organization).where(Organization.slug == "varga"))).scalar_one()
        refs = await pin(session, org.id, project["slug"], ["architecture"])
    return [ref.model_dump(mode="json") for ref in refs]


async def test_a_run_reads_exactly_the_documents_it_was_prepared_with(
    client: AsyncClient, project: dict[str, Any]
) -> None:
    from choregos_api.knowledge.conversion import conversions

    await client.post(BASE, json={"slug": "architecture", "name": "Architecture framework"})
    deck = fx.pptx([("Network zoning principles", "Every zone has a firewall")])
    for name, data in (("zoning.pptx", deck), ("Zoning.md", b"# Zoning notes\n")):
        assert (
            await client.post(f"{BASE}/architecture/documents", files={"file": (name, data)})
        ).status_code == 202
    await conversions.drain()
    pinned = await _pinned(project)
    # Uploaded after the run was prepared: not served to it.
    await client.post(f"{BASE}/architecture/documents", files={"file": ("late.md", b"# Late\n")})
    await conversions.drain()

    token = await _run(project, pinned)
    client.cookies.clear()
    read = await client.get(
        "/api/v1/internal/runs/run-kb/knowledge", headers={"Authorization": f"Bearer {token}"}
    )
    assert read.status_code == 200, read.text
    (base,) = read.json()
    paths = [d["path"] for d in base["documents"]]
    assert sorted(paths) == ["zoning-2.md", "zoning.md"], "two files named alike stay apart"
    assert "late.md" not in paths
    assert base["digest"] == pinned[0]["digest"] == knowledge_digest(base["documents"])
    assert any("# Network zoning principles" in d["markdown"] for d in base["documents"])


async def test_a_base_no_longer_served_to_the_project_is_not_delivered(
    client: AsyncClient, project: dict[str, Any]
) -> None:
    await client.post(BASE, json={"slug": "architecture", "name": "Architecture framework"})
    pinned = await _pinned(project)
    await client.patch(f"{BASE}/architecture", json={"status": "archived"})
    token = await _run(project, pinned)
    client.cookies.clear()
    read = await client.get(
        "/api/v1/internal/runs/run-kb/knowledge", headers={"Authorization": f"Bearer {token}"}
    )
    assert read.status_code == 404


async def test_an_agent_names_its_knowledge_bases_without_changing_older_digests(
    client: AsyncClient, admin: str
) -> None:
    from choregos_api.schemas import AgentSpec
    from choregos_api.services.agents import empreinte

    before = empreinte(AgentSpec(instructions="x"))
    assert before == empreinte(AgentSpec(instructions="x", knowledge=[])), "an empty list changes nothing"
    assert before != empreinte(AgentSpec.model_validate({"instructions": "x", "knowledge": [{"slug": "kb"}]}))
    created = await client.post(
        "/api/v1/orgs/varga/agents",
        json={
            "slug": "architect",
            "display_name": "Architect",
            "spec": {"knowledge": [{"slug": "architecture"}]},
        },
    )
    assert created.status_code == 201, created.text


# -- Spreadsheets that read as figures, decks that keep their notes (S42-09c) ------------------------


def _with_cached_formula(data: bytes, cell: str, cached: str) -> bytes:
    """Writes the value Excel would have saved beside a formula (openpyxl writes formulas without one)."""
    import io
    import re
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            body = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                pattern = re.compile(
                    rb'(<c r="'
                    + cell.encode()
                    + rb'"[^>]*>\s*<f>[^<]*</f>)\s*<v\s*/>|(<c r="'
                    + cell.encode()
                    + rb'"[^>]*>\s*<f>[^<]*</f>)\s*<v></v>'
                )
                body, count = pattern.subn(
                    lambda m: (m.group(1) or m.group(2)) + b"<v>" + cached.encode() + b"</v>", body
                )
                assert count == 1, body[:2000]
            target.writestr(item, body)
    return out.getvalue()


def _workbook() -> bytes:
    import datetime
    import io

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Revenue"
    sheet.append(["Entity", "Revenue", "Margin", "Closed", None, "Note"])
    sheet.append(["Diametral Group", 1234567.891, 0.1234, datetime.date(2025, 3, 1), None, "a | b"])
    sheet.append(["Doubled", "=B2*2", None, None, None, None])
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"].number_format = "0.0%"
    sheet["D2"].number_format = "dd/mm/yyyy"
    sheet["B3"].number_format = "#,##0"
    book.create_sheet("Costs").append(["Item", "Amount"])
    book["Costs"].append(["Rent", 1200])
    secret = book.create_sheet("Scratch")
    secret.append(["draft"])
    secret.sheet_state = "hidden"
    out = io.BytesIO()
    book.save(out)
    return _with_cached_formula(out.getvalue(), "B3", "2469135.782")


def test_a_workbook_reads_as_figures_not_as_a_dataframe() -> None:
    from choregos_api.knowledge.convert import to_markdown

    markdown, _ = to_markdown(_workbook(), ".xlsx")
    assert "NaN" not in markdown and "Unnamed" not in markdown
    assert [line for line in markdown.splitlines() if line.startswith("## ")] == [
        "## Revenue",
        "## Costs",
        "## Scratch (hidden)",
    ]
    # Numbers as the sheet shows them; the formula gives the value Excel saved; a pipe is escaped.
    assert "| Diametral Group | 1,234,567.89 | 12.3% | 2025-03-01 |  | a \\| b |" in markdown
    assert "| Doubled | 2,469,136 |" in markdown
    assert "| Rent | 1200 |" in markdown


def test_a_long_sheet_is_capped_and_says_so() -> None:
    from choregos_api.knowledge.convert import to_markdown

    rows: list[list[object]] = [["id", "value"]] + [[i, i * 2] for i in range(10_000)]
    markdown, _ = to_markdown(fx.xlsx({"Big": rows}), ".xlsx", 2000)
    assert markdown.count("\n| ") == 2000 + 1  # the header, its separator and 1999 rows
    assert "… 8001 more rows not shown" in markdown
    assert "| 9999 |" not in markdown


def test_a_workbook_openpyxl_cannot_read_falls_back_to_markitdown(monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from choregos_api.knowledge import convert, xlsx

    def refuse(*_args: object, **_kwargs: object) -> str:
        raise ValueError("unreadable")

    monkeypatch.setattr(xlsx, "xlsx_to_markdown_with_warnings", refuse)
    markdown, _ = convert.to_markdown(fx.xlsx({"Sheet1": [["a", "b"], [1, 2]]}), ".xlsx")
    assert "## Sheet1" in markdown and "| a | b |" in markdown.replace("  ", " ")


def test_a_deck_keeps_its_speaker_notes_tables_and_chart_text() -> None:
    from choregos_api.knowledge.convert import to_markdown

    deck = fx.pptx(
        [("Group strategy", "Three pillars")],
        table=[["Region", "Share"], ["Nordics", "41%"]],
        notes={0: "Say that the Nordics carry the plan"},
        chart=(["2024", "2025"], {"Revenue": [10.0, 12.5]}),
    )
    markdown, _ = to_markdown(deck, ".pptx")
    assert "Say that the Nordics carry the plan" in markdown
    assert "Nordics" in markdown and "41%" in markdown
    assert "2025" in markdown and "12.5" in markdown


async def test_the_row_cap_reaches_the_child_process_by_argv() -> None:
    """The child's environment is bare by design: the cap is an argument, not a variable."""
    from choregos_api.knowledge.conversion import run_converter

    rows: list[list[object]] = [["id"]] + [[i] for i in range(50)]
    answer = await run_converter(
        fx.xlsx({"Many": rows}), ".xlsx", timeout_s=60, max_bytes=1024 * 1024, rows_per_sheet=10
    )
    assert "… 41 more rows not shown" in answer["markdown"]


def test_a_cut_sheet_is_named_in_the_warnings_and_a_short_one_is_not() -> None:
    """S42-09e: the people who read the document's page are told what the agents are told."""
    from choregos_api.knowledge.convert import convert

    rows: list[list[object]] = [["id"]] + [[i] for i in range(30)]
    book = fx.xlsx({"Big": rows, "Small": [["a"], [1]]})
    markdown, _, warnings = convert(book, ".xlsx", 10)
    assert warnings == ['Sheet "Big": 21 more rows not shown (the first 10 are kept)']
    assert "… 21 more rows not shown" in markdown
    assert convert(book, ".xlsx", 100)[2] == []
    assert convert(b"# Notes\n", ".md")[2] == []


async def test_the_child_process_hands_its_warnings_to_the_parent() -> None:
    from choregos_api.knowledge.conversion import run_converter

    rows: list[list[object]] = [["id"]] + [[i] for i in range(50)]
    answer = await run_converter(
        fx.xlsx({"Many": rows}), ".xlsx", timeout_s=60, max_bytes=1024 * 1024, rows_per_sheet=10
    )
    assert answer["warnings"] == ['Sheet "Many": 41 more rows not shown (the first 10 are kept)']
