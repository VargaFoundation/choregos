"""Small documents built at test time: nothing binary is committed (ADR 0051 tests)."""

from __future__ import annotations

import io
import zipfile


def pptx(slides: list[tuple[str, str]], table: list[list[str]] | None = None) -> bytes:
    """A deck: one slide per `(title, body)`, and a last slide holding `table` if given."""
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    for title, body in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    if table:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Table"
        grid = slide.shapes.add_table(
            len(table), len(table[0]), Inches(1), Inches(2), Inches(6), Inches(1)
        ).table
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                grid.cell(i, j).text = cell
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def docx(sections: list[tuple[int, str, str]], table: list[list[str]] | None = None) -> bytes:
    """A Word document: `(level, heading, paragraph)` per section, then `table` if given."""
    import docx as python_docx

    document = python_docx.Document()
    for level, heading, paragraph in sections:
        document.add_heading(heading, level)
        if paragraph:
            document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                grid.cell(i, j).text = cell
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = book.active if first else book.create_sheet()
        first = False
        assert sheet is not None
        sheet.title = name
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pdf(lines: list[str]) -> bytes:
    """A one-page PDF with `lines` of Helvetica text, written by hand (no PDF library needed)."""
    text = (
        "BT /F1 12 Tf 72 720 Td 14 TL "
        + " ".join(
            "(" + line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)") + ") '" for line in lines
        )
        + " ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(text)).encode() + b" >>\nstream\n" + text.encode() + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def broken_office() -> bytes:
    """A zip that claims to be an Office document and is not one: the converter fails on it."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "not xml at all <<<")
    return out.getvalue()
