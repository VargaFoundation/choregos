# SPDX-License-Identifier: Apache-2.0
"""A workbook as markdown an agent can read as figures (S42-09c).

MarkItDown goes through pandas: empty cells become `NaN`, empty headers `Unnamed: 3`, a percentage is
`0.1234`, a date a timestamp, and a 100 000-row sheet is converted whole. This reads the workbook
directly (openpyxl, read-only, cached values): one `## <sheet>` per sheet, empty cells empty, numbers
shown the way the sheet shows them (percentages, thousands, decimals, dates), and at most
`rows_per_sheet` rows per sheet, the rest counted. A formula gives the value Excel last saved.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime, time
from typing import Any

DEFAULT_ROWS_PER_SHEET = 2000
_QUOTED_OR_BRACKETED = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.')
_THOUSANDS = re.compile(r"[#0?],[#0?]")


def _decimals(section: str) -> tuple[int, int]:
    """(mandatory, optional) decimal places of a number format section."""
    if "." not in section:
        return 0, 0
    run = re.match(r"[0#?]*", section.split(".", 1)[1])
    digits = run.group(0) if run else ""
    return digits.count("0") + digits.count("?"), digits.count("#")


def _general(value: float) -> str:
    if float(value).is_integer() and abs(value) < 1e15:
        return str(int(value))
    text = f"{value:.12g}"
    return f"{value:f}".rstrip("0").rstrip(".") if "e" in text else text


def _fixed(value: float, mandatory: int, optional: int) -> str:
    text = f"{value:.{mandatory + optional}f}"
    if optional and "." in text:
        head, tail = text.split(".")
        tail = tail.rstrip("0").ljust(mandatory, "0")
        text = f"{head}.{tail}" if tail else head
    return text


def number_text(value: float, number_format: str) -> str:
    """A number as its cell format shows it: percent, thousands separator, decimals; else plainly."""
    if isinstance(value, int) and abs(value) >= 10**15:
        return str(value)
    section = _QUOTED_OR_BRACKETED.sub("", (number_format or "General").split(";")[0])
    if section.strip().lower() in {"general", "@", ""} or re.search(r"e[+-]", section, re.I):
        return _general(value)
    mandatory, optional = _decimals(section)
    if "%" in section:
        return _fixed(value * 100, mandatory, optional) + "%"
    if _THOUSANDS.search(section):
        text = _fixed(value, mandatory, optional)
        head, _, tail = text.partition(".")
        grouped = f"{int(head):,}" if head.lstrip("-").isdigit() else head
        if head.startswith("-") and grouped == "0":
            grouped = "-0"
        return f"{grouped}.{tail}" if tail else grouped
    if "0" in section:
        return _fixed(value, mandatory, optional)
    return _general(value)


def cell_text(value: Any, number_format: str = "General") -> str:
    """One cell as a table cell: empty stays empty, a pipe or a line break cannot break the table."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, int | float):
        return number_text(value, number_format)
    return re.sub(r"\s*\n\s*", " ", str(value)).replace("|", "\\|").strip()


def _table(rows: list[list[str]]) -> str:
    width = max(
        (len(row) - next((i for i, c in enumerate(reversed(row)) if c), len(row)) for row in rows), default=0
    )
    if not width:
        return ""
    lines = []
    for index, row in enumerate(rows):
        cells = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("|" + " --- |" * width)
    return "\n".join(lines)


def xlsx_to_markdown(data: bytes, rows_per_sheet: int = DEFAULT_ROWS_PER_SHEET) -> str:
    """The markdown of a workbook; raises whatever openpyxl raises on a file it cannot read."""
    return xlsx_to_markdown_with_warnings(data, rows_per_sheet)[0]


def xlsx_to_markdown_with_warnings(
    data: bytes, rows_per_sheet: int = DEFAULT_ROWS_PER_SHEET
) -> tuple[str, list[str]]:
    """The markdown of a workbook and one sentence per sheet cut at the row cap (S42-09e), for the
    people who look at the document: the markdown already says it to the agents."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    warnings: list[str] = []
    try:
        for sheet in book.worksheets:
            hidden = getattr(sheet, "sheet_state", "visible") != "visible"
            parts.append(f"## {sheet.title}{' (hidden)' if hidden else ''}")
            shown: list[list[str]] = []
            more = 0
            for row in sheet.iter_rows():
                texts = [cell_text(c.value, getattr(c, "number_format", "General")) for c in row]
                if not any(texts):
                    continue  # an empty row says nothing, however many there are
                if len(shown) < rows_per_sheet:
                    shown.append(texts)
                else:
                    more += 1
            table = _table(shown)
            parts.append(table if table else "(empty)")
            if more:
                parts.append(f"… {more} more rows not shown")
                warnings.append(
                    f'Sheet "{sheet.title}": {more} more rows not shown (the first {rows_per_sheet} are kept)'
                )
    finally:
        book.close()
    return "\n\n".join(parts) + "\n", warnings
